from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\gokul\Desktop\PROJECTS\climatetwin\ClimateTwin_IEEE_Paper_Draft.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)

styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'; styles['Normal'].font.size = Pt(10)
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
styles['Normal'].paragraph_format.space_after = Pt(0)
styles['Normal'].paragraph_format.line_spacing = 1.0
for name, size in [('Title',15),('Heading 1',10),('Heading 2',9)]:
    st=styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def shade(cell, fill='D9E2F3'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=tblPr.first_child_found_in('w:tblBorders')
    if b is None: b=OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'B7C3D0'); b.append(el)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E79'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(7.5)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(cells[i],'F3F6F9')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(7.3)
    doc.add_paragraph()
    return t
def para(text='', boldlead=None):
    p=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def h(text, level=1):
    p=doc.add_paragraph(text, style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; return p

title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('ClimateTwin: An Uncertainty-Aware Digital-Twin Prototype for Multi-Hazard Urban Risk Monitoring in Chennai')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Gokul R. and Vishvanth S.\nSchool of Computer Science and Engineering, Vellore Institute of Technology, Chennai, India'); r.font.size=Pt(9)

h('Abstract',1)
para('Coastal cities face compound climate hazards that can disrupt interconnected infrastructure systems. Existing urban digital-twin studies often emphasize isolated hazards, static risk assessment, or visualization, while limited work evaluates uncertainty, sensing failures, infrastructure impacts, and warning value together. This paper presents ClimateTwin, a research prototype for uncertainty-aware urban hazard monitoring in Chennai, India. The framework integrates externally sourced meteorological observations, historical reservoir data, satellite precipitation, terrain information, sensor-quality monitoring, data assimilation, probabilistic forecasting, compound-risk analysis, and impact-oriented warning logic. The implementation combines a FastAPI service layer, MQTT-compatible telemetry abstractions, Ensemble Kalman Filtering, graph-based forecasting, uncertainty quantification, and a web decision-support interface. Real environmental data are explicitly separated from proxy, synthetic, and simulated data. The evaluation protocol uses chronological and event-level validation, calibration metrics, sensor-fault experiments, and decision-value analysis. The current study is a prototype validation contribution rather than a field-validated municipal deployment because complete infrastructure topology, operational impact labels, and operator trials are unavailable. The work provides a reproducible methodology for connecting environmental observations to transparent, uncertainty-aware urban risk warnings.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; run=p.add_run('Index Terms—'); run.bold=True; p.add_run('digital twin, urban resilience, multi-hazard early warning, Chennai, data assimilation, Ensemble Kalman Filter, graph neural network, probabilistic forecasting, uncertainty quantification.')

# body in two columns
body=doc.add_section(WD_SECTION.CONTINUOUS); body.page_width=Inches(8.5); body.page_height=Inches(11); body.top_margin=Inches(0.7); body.bottom_margin=Inches(0.7); body.left_margin=Inches(0.7); body.right_margin=Inches(0.7)
sectPr=body._sectPr; cols=sectPr.xpath('./w:cols')[0] if sectPr.xpath('./w:cols') else OxmlElement('w:cols'); cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'360');
if not sectPr.xpath('./w:cols'): sectPr.append(cols)

h('I. INTRODUCTION')
para('Urban climate hazards increasingly occur as interacting events. Heavy rainfall may coincide with saturated reservoirs, drainage limitations, road disruption, and reduced emergency access. Monitoring systems organized by hazard or agency provide limited support for understanding these interactions and often omit calibrated uncertainty.')
para('Digital twins provide an integration layer for observations, models, simulations, and decision support. However, recent reviews identify unresolved problems in interoperability, real-time synchronization, multi-hazard coupling, infrastructure dependency representation, uncertainty calibration, governance, and operational validation. ClimateTwin addresses a narrower and testable question: can uncertainty-aware state estimation and probabilistic forecasting improve event-level urban hazard warnings when observations are sparse, faulty, or shifted?')
para('The contributions are: 1) a provenance-aware urban risk-twin architecture; 2) comparison of nudging and EnKF state estimation; 3) uncertainty-aware temporal and graph forecasting; 4) compound-risk and impact-warning logic; 5) chronological, event-level, and fault-robustness evaluation; and 6) explicit separation of real, proxy, synthetic, and simulated evidence.')

h('II. RELATED WORK')
h('A. Urban Digital Twins',2); para('Urban digital twins combine city models, sensor streams, simulation, and decision support. Recent studies show broad potential but also report weak interoperability, incomplete operational validation, and limited evidence of transfer across cities. ClimateTwin therefore treats the twin as a research prototype whose synchronization and decision value must be measured rather than assumed.')
h('B. Multi-Hazard and Impact-Based Warning',2); para('International multi-hazard warning guidance emphasizes risk knowledge, monitoring and forecasting, warning dissemination, and preparedness to respond. Hazard intensity alone is insufficient: warnings should describe expected impact, exposed assets, vulnerability, uncertainty, and available response time.')
h('C. Compound Risk and Graph Forecasting',2); para('Compound-hazard studies and graph-based forecasting methods have already established interaction modelling and spatiotemporal graph learning as active research areas. The remaining opportunity is to test whether such methods add value when the physical network is incomplete and when predictions must remain calibrated during extreme events.')
h('D. Research Gap',2); para('The central gap addressed by this work is the lack of a reproducible workflow connecting real observations to state estimation, probabilistic hazard forecasting, infrastructure-impact reasoning, and transparent warnings. The paper does not claim to solve all urban digital-twin gaps; incomplete infrastructure topology, real operator validation, equity analysis, and real cross-city transfer remain limitations.')

h('III. SYSTEM ARCHITECTURE')
para('ClimateTwin contains five logical layers: data acquisition and provenance; edge and telemetry ingestion; state estimation and assimilation; hazard, compound-risk, and impact forecasting; and warning visualization with decision support. Each response records source, retrieval time, data status, model version, uncertainty, and fallback status. The platform includes a FastAPI backend, React dashboard, MQTT-compatible ingestion abstractions, persistent storage, replay, model monitoring, and audit-oriented configuration.')
table(['Layer','Primary function','Evidence status'],[['Data','Open-Meteo, IMERG, NOAA, reservoir, DEM, OSM','Real or historical; provenance required'],['State','Nudging and EnKF assimilation','Prototype; EnKF experimental'],['Forecast','RF, LSTM, GNN-LSTM, quantiles','Measured only where rerun'],['Risk','Compound interactions and cascades','Partly rule-based/simulated'],['Decision','Alerts, replay, expected utility','Simulation unless operator data exist']])

h('IV. DATA SOURCES AND PROVENANCE')
para('The system uses current Open-Meteo weather observations, NASA IMERG precipitation, NOAA observations, historical CMWSSB reservoir records, Copernicus DEM, OpenStreetMap roads/waterways/buildings, and NASA FIRMS hotspots when configured. These sources are real external data, but their presence does not make the downstream risk probability field-validated.')
para('Synthetic or simulated elements include fallback weather data when external APIs fail, simulated pump and grid telemetry, procedurally generated topology, synthetic city-transfer distributions, injected sensor faults, scenario replay values, and simulated warning costs. The dashboard must display whether each value is LIVE, HISTORICAL, PROXY, SIMULATED, or FALLBACK.')
table(['Evidence class','Examples','Permitted claim'],[['Real/live','Open-Meteo, FIRMS, IMERG, NOAA','Externally sourced observation'],['Historical real','CMWSSB, IFI events','Backtest input with provenance'],['Proxy','OSM roads/waterways for missing networks','Geospatial proxy'],['Synthetic','Generated topology or city shift','Methodological simulation'],['Rule-based','Compound score and thresholds','Prototype decision logic']])

h('V. METHODOLOGY')
h('A. Event Definitions and Splits',2); para('Flood, heatwave, extreme rainfall, reservoir stress, and infrastructure-impact labels must be defined before model training. Evaluation uses chronological training, validation, and untouched test periods, with event-level and leave-one-event-out analyses. Feature construction, normalization, thresholds, and calibration must not use future observations.')
h('B. Data Assimilation',2); para('The EnKF maintains an ensemble of reservoir states and updates the ensemble using available observations. Missing observations are masked, while stale, biased, or stuck observations receive inflated observation error. The current process model is simplified persistence with process noise; it is not a complete hydrodynamic simulator.')
h('C. Probabilistic Forecasting',2); para('The forecasting layer compares persistence, threshold, logistic regression, random forest, gradient boosting, LSTM, TCN, GNN, and GNN-LSTM models. Quantile outputs are evaluated using pinball loss, coverage, interval width, weighted interval score, and CRPS. A graph model is only interpreted as physically meaningful when its topology is provenance-tracked.')
h('D. Compound Risk and Impact Warnings',2); para('Compound risk combines hazard components and interaction terms. Because current interactions are partly expert-defined, the output must be labelled as a risk index unless fitted and calibrated against real impact labels. Impact warnings incorporate probability, uncertainty, response time, vulnerability, and asset criticality.')

h('VI. EXPERIMENTAL DESIGN')
table(['Question','Comparison','Primary metrics'],[['RQ1: Assimilation','None vs. nudging vs. EnKF','RMSE, MAE, coverage, runtime'],['RQ2: Forecasting','LSTM vs. GNN-LSTM','RMSE, MAE, CRPS, WIS'],['RQ3: Calibration','Normal vs. extreme events','Coverage, sharpness, reliability'],['RQ4: Compound risk','Additive vs. interacting','Brier, log loss, POD, FAR'],['RQ5: Operations','Threshold vs. utility warning','Lead time, regret, net benefit']])
para('Results must include positive-event counts, confusion matrices, event-level confidence intervals, and exact data-status labels. If a model collapses to an all-negative predictor or if a metric is undefined, the result must remain visible and be discussed rather than replaced with a target value.')

h('VII. CURRENT EVIDENCE AND LIMITATIONS')
para('The repository contains a substantial research prototype: external data ingestion, historical datasets, data-quality checks, chronological-split tooling, EnKF code, GNN-LSTM code, uncertainty modules, compound-risk logic, sensor-fault experiments, GIS processing, replay, dashboard components, tests, and reproducibility documentation.')
para('The evidence does not yet support claims of a fully operational municipal digital twin. Public GIS data provide roads, waterways, buildings, and terrain, but verified pumping-station and substation dependency data are missing. Infrastructure-impact labels are limited, cross-city transfer is synthetic, fault studies are simulated, and operator decision experiments are not field studies. Open-Meteo availability measures API/data-source health, not physical sensor health. The displayed risk percentage is a prototype model output and must not be presented as a field-calibrated probability.')

h('VIII. REQUIRED RESULTS TABLES')
table(['Model','RMSE','F1','POD','FAR','Lead time'],[['Persistence','[TODO]','[TODO]','[TODO]','[TODO]','[TODO]'],['Random Forest','[TODO]','[TODO]','[TODO]','[TODO]','[TODO]'],['LSTM','[TODO]','[TODO]','[TODO]','[TODO]','[TODO]'],['GNN-LSTM','[TODO]','[TODO]','[TODO]','[TODO]','[TODO]']])
table(['Assimilation','RMSE','P90 coverage','Runtime'],[['None','[TODO]','[TODO]','[TODO]'],['Nudging','[TODO]','[TODO]','[TODO]'],['EnKF','[TODO]','[TODO]','[TODO]']])

h('IX. DISCUSSION')
para('The scientific value of ClimateTwin depends on disciplined evidence boundaries. A negative GNN result may indicate that incomplete topology does not support graph learning. Poor calibration during extreme events may reveal exchangeability failure in conformal prediction. High false-alarm rates under sensor faults may motivate uncertainty-aware alert softening. These are publishable findings when measured on fixed test events and reported with limitations.')
para('The most important next validation is one real hazard-to-impact endpoint, such as documented road closure, waterlogging, reservoir overflow, or service disruption. This would move the work from environmental hazard monitoring toward infrastructure-impact early warning.')

h('X. CONCLUSION')
para('ClimateTwin provides a reproducible prototype for integrating environmental observations, state estimation, probabilistic forecasting, compound-risk analysis, and uncertainty-aware warnings for Chennai. Its strongest defensible contribution is not the combination of familiar technologies alone, but the explicit evaluation of how uncertainty, sensing failures, incomplete topology, and evidence provenance affect an urban risk-twin workflow. Until real infrastructure-impact labels, complete dependency topology, and operational stakeholder validation are obtained, the system should be described as prototype validation ready rather than a fully validated municipal digital twin.')

h('REFERENCES')
refs=[
'[1] E. Macatulad and F. Biljecki, “Continuing from the Sendai Framework midterm: Opportunities for urban digital twins in disaster risk management,” International Journal of Disaster Risk Reduction, vol. 102, 104310, 2024.',
'[2] M. Zhu and J. Jin, “Data-driven urban digital twins and critical infrastructure under climate change: A review of frameworks and applications,” Urban Planning, vol. 10, 10109, 2025.',
'[3] S. Ghaffarian, “Rethinking digital twin: Introducing digital risk twin for disaster risk management,” npj Natural Hazards, vol. 2, no. 79, 2025.',
'[4] UNDRR, “Global Status of Multi-Hazard Early Warning Systems 2024,” United Nations Office for Disaster Risk Reduction, 2024.',
'[5] WMO, “Handbook on the Use of Risk Knowledge for Multi-Hazard Early Warning Systems,” World Meteorological Organization, 2024.',
'[6] J. Zscheischler et al., “A typology of compound weather and climate events,” Nature Reviews Earth & Environment, vol. 1, pp. 333–347, 2020.',
'[7] G. Evensen, “The Ensemble Kalman Filter: Theoretical formulation and practical implementation,” Ocean Dynamics, vol. 53, pp. 343–367, 2003.',
'[8] T. N. Kipf and M. Welling, “Semi-supervised classification with graph convolutional networks,” Proc. ICLR, 2017.',
'[9] R. Koenker and G. Bassett, “Regression quantiles,” Econometrica, vol. 46, pp. 33–50, 1978.',
'[10] H. Omrany et al., “The uptake of urban digital twins in the built environment,” Computational Urban Science, vol. 5, no. 20, 2025.',
'[11] ClimateTwin Repository, “ClimateTwin: Urban Multi-Hazard Early Warning,” GitHub, 2026.'
]
for ref in refs: para(ref)

doc.core_properties.title='ClimateTwin IEEE Paper Draft'; doc.core_properties.subject='Uncertainty-aware urban climate digital twin'; doc.core_properties.author='Gokul R. and Vishvanth S.'
doc.save(OUT)
print(OUT)
