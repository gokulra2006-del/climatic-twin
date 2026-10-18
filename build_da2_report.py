from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'/'docx'/'ClimateTwin_DA2_Project_Report.docx'

def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name='Aptos'; run._element.rPr.rFonts.set(qn('w:ascii'),'Aptos')
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor(*color)

def shade(cell, fill):
    tc=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); tc.append(e)

def cell_text(cell,text,header=False):
    cell.text=text; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header: shade(cell,'1F4E79')
    for p in cell.paragraphs:
        p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        for r in p.runs: set_font(r,8.5,header,color=(255,255,255) if header else (0,0,0))

def table(doc, caption, headers, rows, widths=None):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(caption);set_font(r,9,True)
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.autofit=False
    for i,h in enumerate(headers): cell_text(t.rows[0].cells[i],h,True)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cell_text(cells[i],v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def addp(doc,text, lead=None):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.08
    if lead:
        r=p.add_run(lead);set_font(r,10.5,True)
    r=p.add_run(text);set_font(r,10.5)
    return p

def heading(doc,text,level=1):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text);set_font(r,13 if level==1 else 11,True)
    return p

def page(doc): doc.add_page_break()

def main():
    d=Document();s=d.sections[0];s.top_margin=Inches(.65);s.bottom_margin=Inches(.6);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
    normal=d.styles['Normal'];normal.font.name='Aptos';normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');normal.font.size=Pt(10.5)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=foot.add_run('ClimateTwin | Digital Assignment 2 Project Report');set_font(r,8,color=(90,90,90))
    # Page 1
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(34);p.paragraph_format.space_after=Pt(10)
    r=p.add_run('ClimateTwin');set_font(r,22,True)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('An Uncertainty Aware Multi Hazard Digital Twin for Climate Risk Monitoring in Chennai');set_font(r,14,True)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(18);r=p.add_run('Digital Assignment 2 Project Report');set_font(r,12)
    table(d,'Digital Assignment 2 Submission Details',['Item','Details'],[
        ('Assignment','Digital Assignment 2 (DA2) - 10 marks'),('Submission deadline','15 October 2026'),
        ('Purpose','Methodology with a clearly defined novel aspect; frontend and backend implementation; results demonstration.'),
        ('Authors','Gokul R. and Vishvanth S.'),('Institution','School of Computer Science and Engineering, Vellore Institute of Technology, Chennai.'),('Study area','Greater Chennai, India.')], [1.55,5.8])
    heading(d,'Abstract')
    addp(d,'ClimateTwin is a full-stack prototype for monitoring climate-related urban hazards in Greater Chennai. It integrates historical reservoir and precipitation records, weather observations, terrain, open geospatial layers, and live weather feeds into a FastAPI backend and React/Vite dashboard. The novel aspect is an evidence-aware warning workflow: every output is treated according to its data provenance and quality state, while a quality-aware assimilation layer estimates state under missing or faulty observations. The prototype supports flood, heat, drought, wildfire, and extreme-rainfall screening, uncertainty displays, historical replay, and infrastructure-risk reasoning. Preliminary tests show that conventional temporal baselines are useful for the available flood label, whereas the present graph model does not yet outperform an LSTM because verified active infrastructure topology is incomplete. The report therefore demonstrates working implementation and measured prototype results without presenting simulated layers as operational evidence.')
    heading(d,'Keywords',2);addp(d,'urban digital twin; multi-hazard warning; Chennai; data assimilation; uncertainty; FastAPI; React; early warning.')
    heading(d,'Assignment Requirements Compliance',2)
    addp(d,'This report documents the required problem statement (Section 1), methodology and clearly defined novel aspect (Section 2), full frontend and backend implementation (Section 3), results demonstration (Section 4), and conclusion with future work (Section 5).')
    page(d)
    # Page 2
    heading(d,'1. Problem Statement and Research Context')
    addp(d,'Chennai experiences interacting climate stresses: intense northeast-monsoon rainfall and low-relief drainage create flood exposure; high heat and humidity affect health; reservoir stress changes water-security margins; and atmospheric instability can produce localized extreme rainfall. Existing public dashboards frequently display separate indicators, but they do not always show data quality, uncertainty, source provenance, or the link between a hazard signal and a potential service impact. This makes it difficult for an operator to determine whether an alert is trustworthy and what action it should support.')
    heading(d,'1.1 Research Gap')
    addp(d,'A dashboard is not automatically a digital twin. A credible urban risk twin needs a traceable observation-to-warning chain: acquisition, quality control, state estimation, forecast uncertainty, impact reasoning, and accountable alert presentation. The important gap addressed by ClimateTwin is how to make this chain explicit when city data are sparse, sensors can fail, and active drainage, pump, power, and service-dependency maps are not publicly complete. Urban digital-twin and multi-hazard-warning literature highlights synchronization, interoperability, uncertainty, governance, and impact-based communication as central open challenges [1]-[4].')
    heading(d,'1.2 Aim and Objectives')
    table(d,'Table 1. Project objectives',['Objective','Implementation outcome'],[
        ('O1','Ingest real historical and live environmental data with timestamp, plausibility, duplicate, and staleness checks.'),
        ('O2','Estimate the evolving risk state using Newtonian nudging and an Ensemble Kalman Filter (EnKF) prototype.'),
        ('O3','Compare persistence, conventional machine learning, LSTM, and graph-temporal approaches on chronological splits.'),
        ('O4','Expose risk, uncertainty, data quality, source status, and historical replay through a responsive web dashboard.'),
        ('O5','Separate real observations, proxy GIS, rule-based logic, and synthetic simulations to avoid misleading claims.')], [1.0,6.35])
    heading(d,'1.3 Novel Contribution',2)
    addp(d,'The project novelty is not a claim that a new machine-learning algorithm is universally superior. Rather, it is the integration of provenance-aware data quality gates with uncertainty-aware warning logic in a multi-hazard city prototype. The system retains negative results and distinguishes live external observations from proxy topology and simulated experiments. This makes the dashboard useful as a research and decision-support prototype while preserving scientific honesty about what has and has not been validated.')
    page(d)
    # Page 3
    heading(d,'2. Methodology')
    heading(d,'2.1 Data and Provenance Model',2)
    addp(d,'The acquisition layer receives Open-Meteo weather data and, when configured, NASA FIRMS hotspot observations. Historical experiments use reservoir records, satellite precipitation, weather archives, Copernicus terrain information, and documented flood-event inventory records. OpenStreetMap roads, waterways, and buildings provide exposure and connectivity proxies. A provenance tag is stored with each data stream so that fallback weather, generated topology, fault injection, transfer-shift experiments, and rule-defined costs remain visibly classified as synthetic or proxy evidence.')
    table(d,'Table 2. Evidence classification',['Evidence class','Permitted interpretation'],[
        ('Live or historical real','Environmental observations and archived records; suitable for the stated data-quality and forecasting tests.'),
        ('Proxy','Incomplete GIS representation used for exposure or connectivity, not verified active infrastructure.'),
        ('Synthetic or simulated','Fault injection, generated topology, and scenario replay; used for method testing only.'),
        ('Rule based','Thresholds and compound-risk weights; interpretable screening logic, not calibrated failure probabilities.')], [1.8,5.55])
    heading(d,'2.2 Quality Aware State Estimation',2)
    addp(d,'Telemetry is validated for schema, freshness, plausibility, duplicate values, and innovation against the background state. The baseline nudging update combines a background estimate with an observation using a quality-dependent weight. The EnKF maintains an ensemble of state estimates and uses the Kalman gain to update the state from available observations. Missing or suspect sensors are masked or assigned a larger observation-error covariance. The current EnKF process model is persistence plus process noise; it is a controlled state-estimation prototype, not a hydraulic flood simulator [5], [6].')
    heading(d,'2.3 Forecasting and Risk Logic',2)
    addp(d,'For a future horizon H, the forecasting layer estimates the probability of a documented event label from rainfall accumulation, meteorology, reservoir level, storage change, and terrain-related covariates. The benchmark includes persistence, logistic regression, random forest, XGBoost, LSTM, and a typed GNN-LSTM surrogate. The latter produces P10, P50, and P90 quantile outputs, with a training penalty that discourages quantile crossing. Compound risk combines normalized hazard indicators and pairwise interaction terms; until it is fitted to observed impact outcomes, it is reported as a risk index rather than a physical probability.')
    page(d)
    # Page 4
    heading(d,'3. Full System Implementation')
    heading(d,'3.1 Backend Implementation',2)
    addp(d,'The backend is implemented in Python using FastAPI. It exposes endpoints for telemetry ingestion, live weather retrieval, historical time-series access, hazard scores, uncertainty bounds, alert state, replay scenarios, and model-health information. SQLAlchemy persistence supports reproducible event and alert records, while MQTT-compatible ingestion supports streaming telemetry. Validation occurs at the API boundary before data are accepted into the analytical pipeline. The service records source status, timestamp, quality flags, and fallbacks so a dashboard user can understand the origin of every displayed value.')
    heading(d,'3.2 Frontend Implementation',2)
    addp(d,'The frontend is implemented using React and Vite. It presents multi-hazard cards, spatial station views, time series, quality telemetry, uncertainty bands, alert explanations, and a historical replay interface. The dashboard is intentionally designed to show both risk level and data confidence. For example, a high risk score from stale or degraded input is visually distinguishable from a score supported by fresh external observations. The interface also allows an operator to inspect individual hazard contributors instead of relying on a single opaque score.')
    heading(d,'3.3 End to End Workflow',2)
    table(d,'Table 3. ClimateTwin system flow',['Stage','Backend responsibility','Dashboard output'],[
        ('Acquire','Fetch live and historical sources; retain source metadata.','Data-source status and last-update time.'),
        ('Validate','Check timestamp, schema, range, staleness, duplicate, and innovation.','Sensor-health and data-quality indicators.'),
        ('Assimilate','Run nudging or EnKF update with quality-aware observation handling.','Estimated state and confidence range.'),
        ('Forecast','Run baseline, LSTM, or graph-temporal inference; calculate quantiles.','Risk level, P10/P50/P90 uncertainty, lead-time view.'),
        ('Explain and replay','Record contributors and replay a historical forcing sequence.','Feature contribution, alert history, and scenario comparison.')], [1.0,3.15,3.2])
    heading(d,'3.4 Reproducibility and Safety',2)
    addp(d,'The project records chronological split boundaries, model configuration, data status, and experimental artifacts. API fallbacks are labelled instead of silently appearing as live data. The design also includes audit-oriented concepts such as model version, source provenance, alert acknowledgement, and override rationale. These features make the prototype suitable for controlled research demonstration; they do not authorize autonomous reservoir or municipal infrastructure actuation.')
    page(d)
    # Page 5
    heading(d,'4. Results Demonstration and Discussion')
    heading(d,'4.1 Preliminary Benchmark Results',2)
    addp(d,'The available chronological flood benchmark shows that simple and temporal models can produce useful preliminary event-label performance. On the recorded experiment artifact, persistence achieved F1 = 0.471, logistic regression F1 = 0.860, random forest F1 = 0.824, XGBoost F1 = 0.834, and LSTM F1 = 0.851. These values are preliminary benchmark estimates for the declared label and split, not evidence of a field-deployed warning system. The spatial GNN baseline remains incomplete because the project lacks verified pump, power, and active drainage dependencies.')
    table(d,'Table 4. Reported benchmark summary',['Model','F1 score','Interpretation'],[
        ('Persistence','0.471','Reference baseline; high false-alarm tendency.'),('Logistic regression','0.860','Best recorded F1 on the available chronological artifact.'),('Random forest','0.824','Useful nonlinear tabular baseline.'),('XGBoost','0.834','Competitive tabular baseline.'),('LSTM','0.851','Strong temporal baseline.'),('Typed GNN-LSTM','Not claimed superior','Topology is incomplete; requires fair validation with verified dependencies.')], [1.55,1.0,4.8])
    heading(d,'4.2 Assimilation and Uncertainty Evidence',2)
    addp(d,'Controlled EnKF experiments show lower state-estimation error than a no-assimilation baseline under selected simulated fault conditions, but the benefit declines when observations are structurally uninformative or many sensors fail together. The uncertainty artifact records Platt-scaled Brier-score reduction from 0.0546 to 0.0275 and nominal 90% interval coverage of 1.00 with mean width 0.7832 on the recorded split. These are encouraging calibration checks, but more independent extreme-event tests, sharpness measures, CRPS/WIS, and complete event-level confidence intervals are required before a journal-level claim.')
    heading(d,'4.3 Demonstrated Value and Limitations',2)
    addp(d,'The working application demonstrates that a full-stack dashboard can carry data-quality metadata through ingestion, state estimation, forecast presentation, and replay. Its strongest present value is transparency: it prevents a prototype score from looking more certain than its evidence permits. Key limitations are incomplete active-infrastructure topology, limited road-closure and service-disruption labels, a non-hydrodynamic EnKF process model, a small independent holdout-event count, and simulated cascade, transfer, and intervention experiments. These limitations are explicit design targets for the next implementation phase.')
    page(d)
    # Page 6
    heading(d,'5. Conclusion and Future Work')
    addp(d,'ClimateTwin demonstrates a complete frontend and backend prototype for multi-hazard urban risk monitoring in Chennai. It combines real environmental and historical data with quality gates, uncertainty-aware state estimation, temporal forecasting, compound-risk screening, historical replay, and an explainable dashboard. The principal methodological contribution is an evidence-aware architecture that keeps real observations, proxy data, rules, and simulations distinguishable at every stage. The measured benchmark supports further work on calibrated warning under data degradation, while the current results do not justify claims of operational municipal control or verified infrastructure-failure prediction.')
    heading(d,'5.1 Next Implementation Priorities',2)
    table(d,'Table 5. Recommended next steps',['Priority','Action','Acceptance evidence'],[
        ('1','Create a versioned impact catalogue for waterlogging, road closure, and service disruption.','Source-linked labels and event-level precision, recall, calibration, and lead time.'),
        ('2','Add DEM-constrained inundation scenarios.','Validation against flood extent, observed waterlogging, or road-impact records.'),
        ('3','Obtain verified drainage, pumps, substations, and service dependencies.','Fair LSTM-versus-GNN ablation on a real typed topology.'),
        ('4','Evaluate EnKF with real sensor-error characterization.','Observed fault patterns, latency, and state-estimation comparison.'),
        ('5','Add explainability, governance, and accessibility evaluation.','Fidelity and stability tests plus stakeholder-reviewed warning language.')], [0.65,3.25,3.45])
    heading(d,'References')
    refs=[
        '[1] E. Macatulad and F. Biljecki, "Continuing from the Sendai Framework midterm: Opportunities for urban digital twins in disaster risk management," International Journal of Disaster Risk Reduction, vol. 102, Art. no. 104310, 2024.',
        '[2] M. Zhu and J. Jin, "Data-driven urban digital twins and critical infrastructure under climate change: A review of frameworks and applications," Urban Planning, vol. 10, Art. no. 10109, 2025.',
        '[3] UNDRR, Global Status of Multi-Hazard Early Warning Systems 2024. Geneva, Switzerland, 2024.',
        '[4] WMO, Handbook on the Use of Risk Knowledge for Multi-Hazard Early Warning Systems. Geneva, Switzerland, 2024.',
        '[5] G. Evensen, "The Ensemble Kalman Filter: Theoretical formulation and practical implementation," Ocean Dynamics, vol. 53, no. 4, pp. 343-367, 2003.',
        '[6] P. L. Houtekamer and F. Zhang, "Review of the ensemble Kalman filter for atmospheric data assimilation," Monthly Weather Review, vol. 144, no. 12, pp. 4489-4532, 2016.',
        '[7] T. N. Kipf and M. Welling, "Semi-supervised classification with graph convolutional networks," in Proc. ICLR, 2017.',
        '[8] Y. Seo et al., "Structured sequence modeling with graph convolutional recurrent networks," in Proc. ICONIP, 2018.' ]
    for x in refs:
        p=d.add_paragraph();p.paragraph_format.left_indent=Inches(.24);p.paragraph_format.first_line_indent=Inches(-.24);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0;r=p.add_run(x);set_font(r,8.3)
    OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)

if __name__=='__main__':main()
