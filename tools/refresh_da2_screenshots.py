"""Replace the DA2 report's old screenshot appendix with supplied current UI evidence."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPORT = Path(r"C:\Users\gokul\Downloads\ClimateTwin_DA2_Project_Report_With_Screenshots.docx")
OUTPUT = Path(r"C:\Users\gokul\Downloads\ClimateTwin_DA2_Project_Report_Final_Screenshots.docx")

FIGURES = [
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-fc52e576-44dd-4594-ac2a-2d6fd2eb1a11.png",
     "Fig. A1. Live System Digital Twin dashboard. The local demonstration shows the online service state, a moderate prototype flood-risk output, input metrics, and a GIS asset-risk overlay."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-be456949-45ad-44f2-b35f-822ccc98dd4e.png",
     "Fig. A2. Real-Time Flood Intelligence overview. The interface presents the current situation, model outlook, primary forecast drivers, historical analogue, uncertainty-aware risk index, and baseline comparison."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-136d0d63-4347-4f69-aa99-8c62e0d62c85.png",
     "Fig. A3. Flood-intelligence explanation view. This screen documents prediction intervals, telemetry percentiles, nearest historical analogues, and ranked feature drivers for the displayed prototype forecast."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-fbf2159f-5be9-4769-afdd-a296fe55356d.png",
     "Fig. A4. Compound Risk Digital Twin. The screen demonstrates cascading-risk monitoring, a zonal vulnerability index, and an interactive what-if scenario sandbox for controlled prototype analysis."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-a640e845-2e66-4fc2-81a5-eddcb6f61cf0.png",
     "Fig. A5. Heatwave Digital Twin. The interface records the displayed heat conditions, apparent temperature, humidity, wind, derived heat index, and refreshable weather-data status."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-866df933-c83a-4ad6-be14-30cfb8cbd8b1.png",
     "Fig. A6. Drought and Water Stress Intelligence. The dashboard combines a prototype drought-threat index with reservoir storage, root-zone moisture, rainfall, and municipal advisory indicators."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-0455b20f-96ac-4bfd-a32f-32db8a78fb9f.png",
     "Fig. A7. Wildfire and Urban Fringe Fire Intelligence. The screen demonstrates the prototype fire-threat index, fire-weather indicators, hotspot summary, and containment-readiness interpretation."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-6c354adc-da3e-4fe0-a311-58f4c9a8c9a1.png",
     "Fig. A8. Extreme Rainfall and Cloudburst Intelligence. The prototype view integrates convective indicators, rainfall rate, drain-surcharge buffer, and pump-pre-positioning guidance."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-348ac5f6-4cc0-447f-8d46-1864d43fc205.png",
     "Fig. A9. Sensor and Telemetry Infrastructure. The administration screen lists the configured prototype sensor stations and their displayed connectivity state."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-12476aaf-ce4d-4870-894d-daf7b8809f37.png",
     "Fig. A10. Machine Learning Model Management. The administration view documents the active prototype model, validation metric, deployment label, and displayed feature-importance ranking."),
    (r"C:\Users\gokul\AppData\Local\Temp\codex-clipboard-0d453e21-ef83-42fc-ad37-8ad1cd7449c8.png",
     "Fig. A11. Operational Audit Logs and Traceability. The interface demonstrates timestamped prototype audit events for telemetry synchronization, assimilation, inference, drills, and actuation workflows."),
]


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def move_before(paragraph, reference_paragraph):
    reference_paragraph._p.addprevious(paragraph._p)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    tag = qn("w:keepNext")
    existing = p_pr.find(tag)
    if value and existing is None:
        p_pr.append(OxmlElement("w:keepNext"))


def add_before_reference(doc, reference, text=None, style=None):
    paragraph = doc.add_paragraph(text or "", style=style)
    move_before(paragraph, reference)
    return paragraph


def main():
    missing = [str(path) for path, _ in FIGURES if not Path(path).is_file()]
    if missing:
        raise FileNotFoundError("Missing supplied screenshot(s):\n" + "\n".join(missing))

    doc = Document(REPORT)
    paragraphs = list(doc.paragraphs)
    appendix_start = next((i for i, p in enumerate(paragraphs) if p.text.strip() == "Appendix A Project Demonstration Screens"), None)
    references_index = next((i for i, p in enumerate(paragraphs) if p.text.strip() == "References"), None)
    if appendix_start is None or references_index is None or appendix_start >= references_index:
        raise RuntimeError("Could not locate the existing screenshot appendix and References heading.")

    # Delete the old appendix heading, images, captions, and separators, retaining References onward.
    for paragraph in paragraphs[appendix_start:references_index]:
        remove_paragraph(paragraph)

    reference = next(p for p in doc.paragraphs if p.text.strip() == "References")
    heading = add_before_reference(doc, reference, "Appendix A Project Demonstration Screens")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.bold = True
        run.font.size = Pt(12)
    set_keep_with_next(heading)

    intro = add_before_reference(
        doc,
        reference,
        "The following screens document the implemented local prototype interfaces used in the project demonstration. Displayed values are interface evidence and must be interpreted with the system's stated prototype and data-source limitations.",
    )
    intro.paragraph_format.space_after = Pt(8)

    for image_path, caption in FIGURES:
        image_paragraph = add_before_reference(doc, reference)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(8)
        image_paragraph.paragraph_format.space_after = Pt(2)
        run = image_paragraph.add_run()
        run.add_picture(image_path, width=Inches(5.9))

        caption_paragraph = add_before_reference(doc, reference, caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_after = Pt(8)
        for run in caption_paragraph.runs:
            run.italic = True
            run.font.size = Pt(9)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(FIGURES)} supplied screenshots.")


if __name__ == "__main__":
    main()
