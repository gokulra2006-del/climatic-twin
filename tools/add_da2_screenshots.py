from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\gokul\Downloads\ClimateTwin_DA2_Project_Report (1).docx")
OUTPUT = Path(r"C:\Users\gokul\Downloads\ClimateTwin_DA2_Project_Report_With_Screenshots.docx")
TEMP = Path(r"C:\Users\gokul\AppData\Local\Temp")

FIGURES = [
    (
        TEMP / "codex-clipboard-08a04ee4-12f6-4e32-a1d6-b399fa7aaad0.png",
        "Fig. A1. Compound Risk Digital Twin interface. The screen presents a prototype cascading-risk assessment and a what-if scenario sandbox for controlled exploration of rainfall, temperature, and reservoir assumptions.",
    ),
    (
        TEMP / "codex-clipboard-ce792005-7563-406f-9460-9cbd3004f239.png",
        "Fig. A2. Flood dashboard demonstration. The interface combines a prototype risk category, uncertainty bounds, spatial asset overlay, and evidence labels for the Chennai study area.",
    ),
    (
        TEMP / "codex-clipboard-56a6d3f8-3d4c-443e-bf46-a0cb125f5f78.png",
        "Fig. A3. Flood-intelligence demonstration view. When the optional backend is unavailable, the system displays a clearly labelled saved prototype scenario rather than presenting it as live telemetry.",
    ),
    (
        TEMP / "codex-clipboard-57ab0b88-81db-4e1a-b5d9-98508e4f6402.png",
        "Fig. A4. Flood-risk explanation view. The page shows calibrated risk, prediction interval, historical comparison, analog periods, and ranked feature drivers for the saved prototype scenario.",
    ),
    (
        TEMP / "codex-clipboard-7a6bc412-5523-4c8e-b971-2fbbfced9f3f.png",
        "Fig. A5. Early Warning Center demonstration. The centre supports nominal monitoring, visible alert state, drill-only emergency scenarios, and event-history presentation.",
    ),
]


def insert_before(paragraph):
    result = paragraph._parent.add_paragraph()
    result._p.getparent().remove(result._p)
    paragraph._p.addprevious(result._p)
    return result


def apply_font(paragraph, size=11, bold=False, italic=False):
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic


def main():
    missing = [str(image) for image, _ in FIGURES if not image.exists()]
    if missing:
        raise FileNotFoundError("Missing screenshot(s): " + "; ".join(missing))

    doc = Document(SOURCE)
    refs = next(p for p in doc.paragraphs if p.text.strip() == "References")

    # Start the visual appendix on a fresh page without disturbing the main report.
    page_break = insert_before(refs)
    page_break.add_run().add_break(WD_BREAK.PAGE)

    heading = insert_before(refs)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.add_run("Appendix A Project Demonstration Screens")
    apply_font(heading, size=13, bold=True)
    heading.paragraph_format.space_after = Pt(8)

    lead = insert_before(refs)
    lead.add_run(
        "The following screens document the implemented frontend workflows. They are presented as prototype evidence: simulated, proxy, and saved demonstration values remain visibly labelled and are not represented as operational municipal alerts."
    )
    apply_font(lead, size=10)
    lead.paragraph_format.space_after = Pt(8)

    for index, (image_path, caption_text) in enumerate(FIGURES):
        if index:
            new_page = insert_before(refs)
            new_page.add_run().add_break(WD_BREAK.PAGE)

        image_paragraph = insert_before(refs)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_after = Pt(4)
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))

        caption = insert_before(refs)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(caption_text)
        apply_font(caption, size=9, italic=True)
        caption.paragraph_format.space_after = Pt(4)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
