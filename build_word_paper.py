from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'ClimateTwin_IEEE_Submission_Safe_Draft.tex'
OUT = ROOT / 'output' / 'docx' / 'ClimateTwin_IEEE_Submission_Safe_Draft.docx'

def clean(s):
    s = s.replace('~', ' ').replace('\\&', '&').replace('--', '-').replace('``', '"').replace("''", '"')
    s = s.replace('\\begin{enumerate}', '').replace('\\end{enumerate}', '').replace('\\item', '')
    s = re.sub(r'\\(emph|textit|textbf|mathrm)\{([^{}]*)\}', r'\2', s)
    s = re.sub(r'\\cite\{[^}]*\}', '', s)
    s = re.sub(r'\\(ref|label|url)\{[^}]*\}', '', s)
    s = re.sub(r'\\%', '%', s)
    s = re.sub(r'\\[a-zA-Z]+(?:\[[^]]*\])?', '', s)
    s = s.replace('{', '').replace('}', '').replace('$', '')
    return re.sub(r'\s+', ' ', s).strip()

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    tcPr.append(borders)

def set_cell(cell, text, header=False):
    cell.text = text; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)
    if header: shade(cell, '1F4E79')
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        for r in p.runs:
            r.font.name='Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r.font.size=Pt(8.5); r.font.bold=header; r.font.color.rgb=RGBColor(255,255,255) if header else RGBColor(0,0,0)

def add_table(doc, caption, headers, rows, widths=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(caption); r.bold=True; r.font.size=Pt(9)
    table=doc.add_table(rows=1, cols=len(headers)); table.style='Table Grid'; table.autofit=False
    for i,h in enumerate(headers): set_cell(table.rows[0].cells[i], h, True)
    for idx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,v in enumerate(row):
            set_cell(cells[i], v, False)
            if idx % 2: shade(cells[i], 'F3F6F9')
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def setup(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.2; normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',17),('Heading 1',13),('Heading 2',11)]:
        st=styles[name]; st.font.name='Aptos Display' if name=='Title' else 'Aptos'; st._element.rPr.rFonts.set(qn('w:ascii'),st.font.name); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(12 if name!='Title' else 0); st.paragraph_format.space_after=Pt(5)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('ClimateTwin IEEE Research Paper Draft'); run.font.size=Pt(8); run.font.color.rgb=RGBColor(90,90,90)

def add_body(doc, text):
    if not text: return
    p=doc.add_paragraph(style='Normal'); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.add_run(text)

def main():
    tex=SOURCE.read_text(encoding='utf-8')
    doc=Document(); setup(doc)
    title=re.search(r'\\title\{(.*?)\}',tex,re.S).group(1); doc.add_paragraph(clean(title),style='Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('Gokul R. and Vishvanth S.\nSchool of Computer Science and Engineering, Vellore Institute of Technology, Chennai, India'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.font.size=Pt(10)
    abs_m=re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}',tex,re.S); doc.add_heading('Abstract',1); add_body(doc,clean(abs_m.group(1)))
    keys=re.search(r'\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}',tex,re.S)
    p=doc.add_paragraph(); p.add_run('Keywords: ').bold=True; p.add_run(clean(keys.group(1)))
    # Remove document wrappers and bibliography; retain prose plus section hierarchy.
    body=tex.split('\\begin{document}',1)[-1]
    body=re.sub(r'(?s)\\maketitle.*?\\end\{IEEEkeywords\}', '', body)
    body=re.sub(r'(?s)\\begin\{table\*?\}.*?\\end\{table\*?\}', '', body)
    body=re.sub(r'(?s)\\begin\{align\*?\}(.*?)\\end\{align\*?\}', r'\\[Equation: \1\\]', body)
    body=re.sub(r'(?s)\\begin\{equation\*?\}(.*?)\\end\{equation\*?\}', r'\\[Equation: \1\\]', body)
    body=body.split('\\begin{thebibliography}',1)[0]
    tokens=re.split(r'(\\(?:sub)*section\{.*?\})',body,flags=re.S)
    for token in tokens:
        match=re.fullmatch(r'\\(sub)*section\{(.*?)\}',token.strip(),re.S)
        if match:
            doc.add_heading(clean(match.group(2)),2 if match.group(1) else 1); continue
        for para in re.split(r'\n\s*\n',token):
            para=para.strip()
            if not para or para.startswith('\\end{document}'): continue
            if para.startswith('\\begin{enumerate}'):
                para=para.replace('\\begin{enumerate}','').replace('\\end{enumerate}','')
                for item in re.split(r'\\item',para):
                    item=clean(item)
                    if item: doc.add_paragraph(item,style='List Bullet')
                continue
            text=clean(para)
            if len(text)>35:
                if text.startswith('[Equation:'):
                    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); r.italic=True; r.font.size=Pt(10)
                else: add_body(doc,text)
    doc.add_heading('Evidence and Feature Maturity Summary',1)
    add_body(doc,'The following matrices convert the prototype architecture into a publication-ready evaluation plan. They summarize what is implemented, what is experimentally assessed, and what evidence must be added before a deployment or impact-forecasting claim can be made.')
    add_table(doc,'Table 1. Evidence provenance used in ClimateTwin.', ['Class','Examples and allowed interpretation'], [
        ('Live external','Open-Meteo weather and optionally NASA FIRMS observations. External observations, not municipal sensor telemetry.'),
        ('Historical real','CMWSSB reservoir records, IMERG precipitation, NOAA weather, Copernicus DEM, and documented historical flood inventories.'),
        ('Proxy','OpenStreetMap roads, waterways, and buildings used as incomplete exposure or connectivity layers.'),
        ('Synthetic or simulated','Fault injection, generated topology, transfer shifts, operator costs, and scenario replay settings.'),
        ('Rule based','Thresholds and compound-risk coefficients not yet fitted to observed impact outcomes.')], [1.5,5.7])
    add_table(doc,'Table 2. Feature maturity and evidence required for stronger claims.', ['Feature','Current status','Evidence required'], [
        ('Live weather ingestion','Implemented','Uptime, latency, provenance, and fallback-rate monitoring.'),
        ('Historical flood benchmark','Preliminary evaluation','Frozen labels, exact event list, independent rerun, and confidence intervals.'),
        ('Nudging and EnKF','Prototype experiment','Real observation-error characterization and state-estimation validation.'),
        ('LSTM and conventional ML','Preliminary benchmark','Fixed threshold policy and an independently observed impact endpoint.'),
        ('GNN-LSTM','Prototype experiment','Verified typed asset dependencies and fair topology ablations.'),
        ('Compound risk and cascades','Rule and simulation prototype','Calibration against observed compound impacts and failure records.'),
        ('Scenario replay and MPC','Proposed decision support','Verified operations, physical constraints, and stakeholder authorization.')], [1.45,1.55,4.2])
    doc.add_heading('References',1)
    bib=re.search(r'(?s)\\begin\{thebibliography\}\{.*?\}(.*?)\\end\{thebibliography\}',tex)
    if bib:
        entries=re.split(r'\\bibitem\{[^}]*\}',bib.group(1))
        for e in entries:
            e=clean(e)
            if e:
                p=doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
                r=p.add_run(e); r.font.size=Pt(8.5)
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)

if __name__=='__main__': main()
